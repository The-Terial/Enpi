from __future__ import annotations

from copy import copy
from datetime import datetime
from math import sqrt
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT = Path("outputs/excel_metrel/Lecturas_con_Metrel_Mi2883_trabajo.xlsx")
OUTPUT = Path("outputs/excel_metrel/Lecturas_con_Metrel_Mi2883_metricas.xlsx")

SCENARIOS = ["013", "012", "011", "010", "009", "008", "007"]
RAW_SCENARIOS = ["011", "010", "009", "008", "007"]
MA_HEADERS = ["Error (Enp - Metrel)", "Error Abs", "Error Cuadrado", "Error porcentual abs"]
VARIABLES = [
    ("Voltaje", "V"),
    ("energia_kwh", "kWh"),
    ("Frecuencia", "Hz"),
    ("Factor_Potencia", ""),
    ("corriente", "A"),
    ("potencia_kw", "kW"),
]


def floor_to_second(value):
    if isinstance(value, datetime):
        return value.replace(microsecond=0)
    return value


def copy_cell_format(src, dst):
    if src.has_style:
        dst._style = copy(src._style)
    if src.number_format:
        dst.number_format = src.number_format
    if src.alignment:
        dst.alignment = copy(src.alignment)
    if src.font:
        dst.font = copy(src.font)
    if src.fill:
        dst.fill = copy(src.fill)
    if src.border:
        dst.border = copy(src.border)


def detect_current_unit(ws):
    headers = [str(ws.cell(1, col).value or "") for col in range(1, ws.max_column + 1)]
    return "mA" if any("I1(Med) [mA]" in h for h in headers) else "A"


def has_direct_power(ws):
    headers = [str(ws.cell(1, col).value or "") for col in range(1, ws.max_column + 1)]
    return any("P1+(Med) [W]" in h for h in headers)


def to_float(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def unscale_if_needed(value):
    value = to_float(value)
    if value is None:
        return None
    return value / 1000000 if abs(value) > 1000 else value


def normalize_pagina(ws):
    headers = ["Voltaje", "energia_kwh", "Frecuencia", "Factor_Potencia", "corriente", "potencia_kw"]
    for idx, header in enumerate(headers, start=12):  # L:Q
        ws.cell(1, idx).value = header
        ws.cell(1, idx).font = Font(bold=True)

    for row in range(2, ws.max_row + 1):
        voltaje = unscale_if_needed(ws.cell(row, 4).value)
        corriente = unscale_if_needed(ws.cell(row, 5).value)
        potencia = unscale_if_needed(ws.cell(row, 6).value)
        energia = unscale_if_needed(ws.cell(row, 7).value)
        frecuencia = unscale_if_needed(ws.cell(row, 10).value)
        fp = unscale_if_needed(ws.cell(row, 11).value)

        ws.cell(row, 12).value = voltaje
        ws.cell(row, 13).value = energia
        ws.cell(row, 14).value = frecuencia
        ws.cell(row, 15).value = fp
        ws.cell(row, 16).value = corriente
        ws.cell(row, 17).value = potencia

        for col in range(12, 18):
            ws.cell(row, col).number_format = "0.000000"


def prepare_raw_sheet(ws, template_ws):
    ws.insert_cols(1, 3)

    ws["A1"] = "Fecha"
    ws["B1"] = "Hora"
    ws["C1"] = "Fecha Hora"
    ws["C1"].comment = Comment('=TEXTO(A2;"dd-mm-yyyy")&"  "&TEXTO(B2;"hh:mm:ss")', "Codex")

    for col in range(1, 4):
        copy_cell_format(template_ws.cell(1, col), ws.cell(1, col))
        ws.column_dimensions[get_column_letter(col)].width = template_ws.column_dimensions[get_column_letter(col)].width

    for row in range(2, ws.max_row + 1):
        ws.cell(row, 1).value = f"=+INT(D{row})"
        ws.cell(row, 2).value = f"=+MOD(D{row},1)"
        ws.cell(row, 3).value = floor_to_second(ws.cell(row, 4).value)
        copy_cell_format(template_ws["A2"], ws.cell(row, 1))
        copy_cell_format(template_ws["B2"], ws.cell(row, 2))
        copy_cell_format(template_ws["C2"], ws.cell(row, 3))


def sheet_layout(ws):
    direct_power = has_direct_power(ws)
    if direct_power:
        enpi_start = 18  # R
        pf_col = "Q"
        direct_power_col = "P"
    else:
        enpi_start = 17  # Q
        pf_col = "P"
        direct_power_col = None

    enpi_cols = {
        "Voltaje": get_column_letter(enpi_start),
        "energia_kwh": get_column_letter(enpi_start + 1),
        "Frecuencia": get_column_letter(enpi_start + 2),
        "Factor_Potencia": get_column_letter(enpi_start + 3),
        "corriente": get_column_letter(enpi_start + 4),
        "potencia_kw": get_column_letter(enpi_start + 5),
    }
    error_start = enpi_start + 6
    hidden_start = error_start + len(VARIABLES) * 4

    return {
        "direct_power": direct_power,
        "current_unit": detect_current_unit(ws),
        "enpi_start": enpi_start,
        "enpi_cols": enpi_cols,
        "error_start": error_start,
        "hidden_start": hidden_start,
        "pf_col": pf_col,
        "direct_power_col": direct_power_col,
    }


def metrel_ref_formula(var_name, row, layout):
    current_expr = f"J{row}/1000" if layout["current_unit"] == "mA" else f"J{row}"
    if var_name == "Voltaje":
        return f"=F{row}"
    if var_name == "Frecuencia":
        return f"=N{row}"
    if var_name == "Factor_Potencia":
        return f"={layout['pf_col']}{row}"
    if var_name == "corriente":
        return f"={current_expr}"
    if var_name == "potencia_kw":
        if layout["direct_power_col"]:
            return f"={layout['direct_power_col']}{row}/1000"
        return f"=F{row}*({current_expr})*{layout['pf_col']}{row}/1000"
    raise ValueError(var_name)


def apply_analysis(ws):
    layout = sheet_layout(ws)
    max_row = ws.max_row
    enpi_start = layout["enpi_start"]

    # ENPI lookup columns. Pagina!L:Q is normalized to numeric values first.
    enpi_headers = ["Voltaje", "energia_kwh", "Frecuencia", "Factor_Potencia", "corriente", "potencia_kw"]
    enpi_formulas = {
        "Voltaje": '=IFERROR(VLOOKUP($C{row},Pagina!$C:$Q,10,0),"")',
        "energia_kwh": '=IFERROR(VLOOKUP($C{row},Pagina!$C:$Q,11,0),"")',
        "Frecuencia": '=IFERROR(VLOOKUP($C{row},Pagina!$C:$Q,12,0),"")',
        "Factor_Potencia": '=IFERROR(VLOOKUP($C{row},Pagina!$C:$Q,13,0),"")',
        "corriente": '=IFERROR(VLOOKUP($C{row},Pagina!$C:$Q,14,0),"")',
        "potencia_kw": '=IFERROR(VLOOKUP($C{row},Pagina!$C:$Q,15,0),"")',
    }
    for idx, header in enumerate(enpi_headers):
        col = enpi_start + idx
        ws.cell(1, col).value = header
        for row in range(2, max_row + 1):
            ws.cell(row, col).value = enpi_formulas[header].format(row=row)

    # Hidden Metrel reference/helper columns.
    hidden_cols = {}
    for idx, (var_name, _unit) in enumerate(VARIABLES):
        col = layout["hidden_start"] + idx
        hidden_cols[var_name] = get_column_letter(col)
        ws.cell(1, col).value = f"Metrel_{var_name}_ref"
        ws.column_dimensions[get_column_letter(col)].hidden = True

    for row in range(2, max_row + 1):
        for var_name, _unit in VARIABLES:
            col_letter = hidden_cols[var_name]
            if var_name == "energia_kwh":
                enpi_energy = layout["enpi_cols"]["energia_kwh"]
                metrel_power = hidden_cols["potencia_kw"]
                if row == 2:
                    ws[f"{col_letter}{row}"] = f'=IF({enpi_energy}{row}="","",{enpi_energy}{row})'
                else:
                    ws[f"{col_letter}{row}"] = (
                        f'=IF({col_letter}{row-1}="","",{col_letter}{row-1}+IFERROR({metrel_power}{row}*($C{row}-$C{row-1})*24,0))'
                    )
            else:
                ws[f"{col_letter}{row}"] = metrel_ref_formula(var_name, row, layout)

    # Visible error blocks.
    metric_ranges = {}
    for v_idx, (var_name, _unit) in enumerate(VARIABLES):
        start_col = layout["error_start"] + v_idx * 4
        enpi_col = layout["enpi_cols"][var_name]
        metrel_col = hidden_cols[var_name]
        block_letters = [get_column_letter(start_col + i) for i in range(4)]
        metric_ranges[var_name] = block_letters

        for i, header in enumerate(MA_HEADERS):
            ws.cell(1, start_col + i).value = f"{var_name} - {header}"

        for row in range(2, max_row + 1):
            err, abserr, sqerr, pcterr = block_letters
            ws[f"{err}{row}"] = f'=IFERROR({enpi_col}{row}-{metrel_col}{row},"")'
            ws[f"{abserr}{row}"] = f'=IF({err}{row}="","",ABS({err}{row}))'
            ws[f"{sqerr}{row}"] = f'=IF({err}{row}="","",{err}{row}*{err}{row})'
            ws[f"{pcterr}{row}"] = f'=IFERROR(ABS({err}{row}/{metrel_col}{row})*100,"")'

    # Formatting and widths.
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col in range(enpi_start, layout["hidden_start"] + len(VARIABLES)):
        c = ws.cell(1, col)
        c.font = Font(bold=True)
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = 16

    for row in range(2, max_row + 1):
        for col in range(enpi_start, layout["hidden_start"] + len(VARIABLES)):
            ws.cell(row, col).number_format = "0.000000"

    return metric_ranges


def create_summary(wb, all_metric_ranges):
    if "Resumen Errores" in wb.sheetnames:
        del wb["Resumen Errores"]
    ws = wb.create_sheet("Resumen Errores", 1)
    headers = ["Escenario", "Variable", "N", "MAE", "RMSE", "MAPE %"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")

    row_out = 2
    for scenario in SCENARIOS:
        src = wb[scenario]
        max_row = src.max_row
        ranges = all_metric_ranges[scenario]
        for var_name, _unit in VARIABLES:
            err, abserr, sqerr, pcterr = ranges[var_name]
            ws.cell(row_out, 1).value = scenario
            ws.cell(row_out, 2).value = var_name
            ws.cell(row_out, 3).value = f'=COUNT(\'{scenario}\'!{abserr}2:{abserr}{max_row})'
            ws.cell(row_out, 4).value = f'=IF(C{row_out}=0,"",AVERAGE(\'{scenario}\'!{abserr}2:{abserr}{max_row}))'
            ws.cell(row_out, 5).value = f'=IF(C{row_out}=0,"",SQRT(AVERAGE(\'{scenario}\'!{sqerr}2:{sqerr}{max_row})))'
            ws.cell(row_out, 6).value = f'=IF(C{row_out}=0,"",AVERAGE(\'{scenario}\'!{pcterr}2:{pcterr}{max_row}))'
            row_out += 1

    for col in range(1, 7):
        ws.column_dimensions[get_column_letter(col)].width = [12, 22, 10, 14, 14, 14][col - 1]
    for row in range(2, row_out):
        for col in range(3, 7):
            ws.cell(row, col).number_format = "0.000000"
    ws.freeze_panes = "A2"


def get_pagina_map(wb):
    ws = wb["Pagina"]
    result = {}
    for row in range(2, ws.max_row + 1):
        key = floor_to_second(ws.cell(row, 3).value)
        if not key:
            continue
        result[key] = {
            "Voltaje": to_float(ws.cell(row, 12).value),
            "energia_kwh": to_float(ws.cell(row, 13).value),
            "Frecuencia": to_float(ws.cell(row, 14).value),
            "Factor_Potencia": to_float(ws.cell(row, 15).value),
            "corriente": to_float(ws.cell(row, 16).value),
            "potencia_kw": to_float(ws.cell(row, 17).value),
        }
    return result


def metrel_refs_for_row(ws, row, layout):
    current = to_float(ws[f"J{row}"].value)
    if current is not None and layout["current_unit"] == "mA":
        current = current / 1000

    voltage = to_float(ws[f"F{row}"].value)
    frequency = to_float(ws[f"N{row}"].value)
    pf = to_float(ws[f"{layout['pf_col']}{row}"].value)

    if layout["direct_power_col"]:
        p_raw = to_float(ws[f"{layout['direct_power_col']}{row}"].value)
        power_kw = None if p_raw is None else p_raw / 1000
    elif voltage is not None and current is not None and pf is not None:
        power_kw = voltage * current * pf / 1000
    else:
        power_kw = None

    return {
        "Voltaje": voltage,
        "Frecuencia": frequency,
        "Factor_Potencia": pf,
        "corriente": current,
        "potencia_kw": power_kw,
    }


def compute_metric_summary(wb):
    pagina = get_pagina_map(wb)
    summary = {}

    for scenario in SCENARIOS:
        ws = wb[scenario]
        layout = sheet_layout(ws)
        errors = {name: [] for name, _unit in VARIABLES}
        pct_errors = {name: [] for name, _unit in VARIABLES}
        metrel_energy = None
        prev_time = None

        for row in range(2, ws.max_row + 1):
            timestamp = floor_to_second(ws.cell(row, 3).value)
            if not isinstance(timestamp, datetime):
                continue

            refs = metrel_refs_for_row(ws, row, layout)
            enpi = pagina.get(timestamp)

            if prev_time is not None and metrel_energy is not None and refs["potencia_kw"] is not None:
                delta_hours = (timestamp - prev_time).total_seconds() / 3600
                if delta_hours >= 0:
                    metrel_energy += refs["potencia_kw"] * delta_hours
            prev_time = timestamp

            if enpi is None:
                continue

            if metrel_energy is None and enpi["energia_kwh"] is not None:
                metrel_energy = enpi["energia_kwh"]

            refs["energia_kwh"] = metrel_energy

            for var_name, _unit in VARIABLES:
                e_val = enpi.get(var_name)
                m_val = refs.get(var_name)
                if e_val is None or m_val is None:
                    continue
                err = e_val - m_val
                errors[var_name].append(err)
                if m_val != 0:
                    pct_errors[var_name].append(abs(err / m_val) * 100)

        summary[scenario] = {}
        for var_name, _unit in VARIABLES:
            err_values = errors[var_name]
            abs_values = [abs(x) for x in err_values]
            sq_values = [x * x for x in err_values]
            summary[scenario][var_name] = {
                "N": len(err_values),
                "MAE": sum(abs_values) / len(abs_values) if abs_values else None,
                "RMSE": sqrt(sum(sq_values) / len(sq_values)) if sq_values else None,
                "MAPE": sum(pct_errors[var_name]) / len(pct_errors[var_name]) if pct_errors[var_name] else None,
            }
    return summary


def write_static_summary(wb, summary):
    ws = wb["Resumen Errores"]
    row_out = 2
    for scenario in SCENARIOS:
        for var_name, _unit in VARIABLES:
            vals = summary[scenario][var_name]
            ws.cell(row_out, 1).value = scenario
            ws.cell(row_out, 2).value = var_name
            ws.cell(row_out, 3).value = vals["N"]
            ws.cell(row_out, 4).value = vals["MAE"]
            ws.cell(row_out, 5).value = vals["RMSE"]
            ws.cell(row_out, 6).value = vals["MAPE"]
            row_out += 1


def main():
    wb = load_workbook(INPUT, data_only=False)
    template = wb["012"]

    normalize_pagina(wb["Pagina"])

    for sheet_name in RAW_SCENARIOS:
        prepare_raw_sheet(wb[sheet_name], template)

    all_metric_ranges = {}
    for sheet_name in SCENARIOS:
        all_metric_ranges[sheet_name] = apply_analysis(wb[sheet_name])

    create_summary(wb, all_metric_ranges)
    write_static_summary(wb, compute_metric_summary(wb))

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
    except Exception:
        pass

    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
