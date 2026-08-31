from __future__ import annotations

from collections import defaultdict
from math import sqrt
from pathlib import Path
from statistics import mean

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT = Path("/Users/andresviverosllabres/Desktop/mediciones.xlsx")
OUTPUT = Path("outputs/excel_metrel/mediciones_con_indicadores.xlsx")


def as_float(value):
    if value is None or value == "" or value == "---":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "."))
    except ValueError:
        return None


def corr_r2(xs, ys):
    if len(xs) < 2:
        return None
    mx = mean(xs)
    my = mean(ys)
    sx = sum((x - mx) ** 2 for x in xs)
    sy = sum((y - my) ** 2 for y in ys)
    if sx == 0 or sy == 0:
        return None
    cov = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    r = cov / sqrt(sx * sy)
    return r * r


def metrics(pairs):
    # pairs are (metrel, enpi)
    clean = [(m, e) for m, e in pairs if m is not None and e is not None]
    if not clean:
        return {
            "N": 0,
            "MAE": None,
            "RMSE": None,
            "MAPE %": None,
            "Bias": None,
            "Max Abs Error": None,
            "R2": None,
        }

    errors = [e - m for m, e in clean]
    abs_errors = [abs(x) for x in errors]
    sq_errors = [x * x for x in errors]
    pct_errors = [abs((e - m) / m) * 100 for m, e in clean if m != 0]
    xs = [m for m, _e in clean]
    ys = [e for _m, e in clean]

    return {
        "N": len(clean),
        "MAE": mean(abs_errors),
        "RMSE": sqrt(mean(sq_errors)),
        "MAPE %": mean(pct_errors) if pct_errors else None,
        "Bias": mean(errors),
        "Max Abs Error": max(abs_errors),
        "R2": corr_r2(xs, ys),
    }


def current_scale_for_scenario(data_wb, scenario):
    if scenario not in data_wb.sheetnames:
        return 0.001
    header = str(data_wb[scenario].cell(1, 7).value or "")
    if "[A]" in header:
        return 1.0
    return 0.001


def row_pairs(ws, row, current_scale):
    metrel_v = as_float(ws.cell(row, 7).value)  # G U1(Med)
    enpi_v = as_float(ws.cell(row, 18).value)  # R

    metrel_i_ma = as_float(ws.cell(row, 11).value)  # K I1(Med) [mA]
    metrel_i = None if metrel_i_ma is None else metrel_i_ma * current_scale
    enpi_i = as_float(ws.cell(row, 19).value)  # S

    metrel_f = as_float(ws.cell(row, 15).value)  # O f(ProAct)
    enpi_f = as_float(ws.cell(row, 22).value)  # V

    metrel_pf = as_float(ws.cell(row, 17).value)  # Q
    enpi_pf = as_float(ws.cell(row, 23).value)  # W

    enpi_p = as_float(ws.cell(row, 20).value)  # T potencia_kw
    metrel_p = None
    if metrel_v is not None and metrel_i is not None and metrel_pf is not None:
        metrel_p = metrel_v * metrel_i * metrel_pf / 1000

    return {
        "Voltaje": (metrel_v, enpi_v),
        "Corriente": (metrel_i, enpi_i),
        "Frecuencia": (metrel_f, enpi_f),
        "Factor_Potencia": (metrel_pf, enpi_pf),
        "Potencia_kW_derivada": (metrel_p, enpi_p),
    }


def build_summary(data_wb, output_wb):
    ws = data_wb["Todos"]
    data = defaultdict(lambda: defaultdict(list))
    all_key = "Todos"

    for row in range(2, ws.max_row + 1):
        scenario = ws.cell(row, 4).value or "Sin escenario"
        scenario = str(scenario)
        current_scale = current_scale_for_scenario(data_wb, scenario)
        pairs = row_pairs(ws, row, current_scale)
        for var, pair in pairs.items():
            data[scenario][var].append(pair)
            data[all_key][var].append(pair)

    if "Indicadores" in output_wb.sheetnames:
        del output_wb["Indicadores"]
    out = output_wb.create_sheet("Indicadores", 0)

    headers = [
        "Escenario",
        "Variable",
        "N",
        "MAE",
        "RMSE",
        "MAPE %",
        "Bias ENPI-Metrel",
        "Max Abs Error",
        "R2",
        "Unidad",
    ]
    out.append(headers)

    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in out[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    units = {
        "Voltaje": "V",
        "Corriente": "A",
        "Frecuencia": "Hz",
        "Factor_Potencia": "",
        "Potencia_kW_derivada": "kW",
    }

    ordered_scenarios = sorted(k for k in data.keys() if k != all_key) + [all_key]
    ordered_vars = ["Voltaje", "Corriente", "Frecuencia", "Factor_Potencia", "Potencia_kW_derivada"]

    for scenario in ordered_scenarios:
        for var in ordered_vars:
            m = metrics(data[scenario][var])
            out.append([
                scenario,
                var,
                m["N"],
                m["MAE"],
                m["RMSE"],
                m["MAPE %"],
                m["Bias"],
                m["Max Abs Error"],
                m["R2"],
                units[var],
            ])

    for col, width in enumerate([14, 24, 10, 14, 14, 14, 18, 16, 12, 10], start=1):
        out.column_dimensions[get_column_letter(col)].width = width

    for row in range(2, out.max_row + 1):
        for col in range(4, 10):
            out.cell(row, col).number_format = "0.000000"

    out.freeze_panes = "A2"
    out.set_printer_settings(paper_size=out.PAPERSIZE_LETTER, orientation=out.ORIENTATION_LANDSCAPE)
    out.sheet_properties.pageSetUpPr.fitToPage = True
    out.page_setup.fitToWidth = 1
    out.page_setup.fitToHeight = 0
    out.page_margins.left = 0.25
    out.page_margins.right = 0.25
    out.page_margins.top = 0.35
    out.page_margins.bottom = 0.35

    note_row = out.max_row + 2
    out.cell(note_row, 1).value = "Notas"
    out.cell(note_row, 1).font = Font(bold=True)
    out.cell(note_row + 1, 1).value = (
        "Potencia_kW_derivada usa Metrel V*I*FP/1000; solo se calcula cuando el factor de potencia Metrel es numérico."
    )
    out.cell(note_row + 2, 1).value = (
        "MAPE omite filas con referencia Metrel igual a cero. Las filas con '---' o errores se excluyen del N."
    )
    out.merge_cells(start_row=note_row + 1, start_column=1, end_row=note_row + 1, end_column=10)
    out.merge_cells(start_row=note_row + 2, start_column=1, end_row=note_row + 2, end_column=10)


def main():
    data_wb = load_workbook(INPUT, data_only=True)
    output_wb = load_workbook(INPUT, data_only=False)
    build_summary(data_wb, output_wb)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
