from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


INPUT = Path("/Users/andresviverosllabres/Desktop/mediciones_corregido.xlsx")
OUTPUT = Path("outputs/excel_metrel/mediciones_corregido_con_calculos.xlsx")


SCENARIOS = ["000", "001", "002", "003", "004", "005", "006", "007", "Todos"]

VARIABLES = [
    {
        "name": "Voltaje",
        "unit": "V",
        "ref": "G",
        "enpi": "R",
        "err": "AU",
        "abs": "AZ",
        "ape": "BE",
        "sq": "BJ",
    },
    {
        "name": "Corriente",
        "unit": "A",
        "ref": "AC",
        "enpi": "S",
        "err": "AV",
        "abs": "BA",
        "ape": "BF",
        "sq": "BK",
    },
    {
        "name": "Frecuencia",
        "unit": "Hz",
        "ref": "O",
        "enpi": "V",
        "err": "AW",
        "abs": "BB",
        "ape": "BG",
        "sq": "BL",
    },
    {
        "name": "Factor_Potencia",
        "unit": "",
        "ref": "Q",
        "enpi": "W",
        "err": "AX",
        "abs": "BC",
        "ape": "BH",
        "sq": "BM",
    },
    {
        "name": "Potencia_kW_derivada",
        "unit": "kW",
        "ref": "AD",
        "enpi": "T",
        "err": "AY",
        "abs": "BD",
        "ape": "BI",
        "sq": "BN",
    },
]


def metric_formula(metric: str, scenario_cell: str, var: dict[str, str], first_row: int, last_row: int) -> str:
    scen_rng = f"'Todos'!$D${first_row}:$D${last_row}"
    ref_rng = f"'Todos'!${var['ref']}${first_row}:${var['ref']}${last_row}"
    enpi_rng = f"'Todos'!${var['enpi']}${first_row}:${var['enpi']}${last_row}"
    err_rng = f"'Todos'!${var['err']}${first_row}:${var['err']}${last_row}"
    abs_rng = f"'Todos'!${var['abs']}${first_row}:${var['abs']}${last_row}"
    ape_rng = f"'Todos'!${var['ape']}${first_row}:${var['ape']}${last_row}"
    sq_rng = f"'Todos'!${var['sq']}${first_row}:${var['sq']}${last_row}"

    if metric == "N":
        return (
            f'=IF({scenario_cell}="Todos",COUNT({abs_rng}),'
            f'COUNTIFS({scen_rng},{scenario_cell},{abs_rng},">=0"))'
        )
    if metric == "MAE":
        return (
            f'=IFERROR(IF({scenario_cell}="Todos",AVERAGE({abs_rng}),'
            f'AVERAGEIFS({abs_rng},{scen_rng},{scenario_cell})),"")'
        )
    if metric == "RMSE":
        return (
            f'=IFERROR(SQRT(IF({scenario_cell}="Todos",AVERAGE({sq_rng}),'
            f'AVERAGEIFS({sq_rng},{scen_rng},{scenario_cell}))),"")'
        )
    if metric == "MAPE":
        return (
            f'=IFERROR(IF({scenario_cell}="Todos",AVERAGE({ape_rng}),'
            f'AVERAGEIFS({ape_rng},{scen_rng},{scenario_cell})),"")'
        )
    if metric == "Bias":
        return (
            f'=IFERROR(IF({scenario_cell}="Todos",AVERAGE({err_rng}),'
            f'AVERAGEIFS({err_rng},{scen_rng},{scenario_cell})),"")'
        )
    if metric == "Max":
        return (
            f'=IFERROR(IF({scenario_cell}="Todos",MAX({abs_rng}),'
            f'MAXIFS({abs_rng},{scen_rng},{scenario_cell})),"")'
        )
    if metric == "R2":
        valid_all = f"ISNUMBER({ref_rng})*ISNUMBER({enpi_rng})"
        valid_scen = f"({scen_rng}={scenario_cell})*ISNUMBER({ref_rng})*ISNUMBER({enpi_rng})"
        return (
            f'=IFERROR(IF({scenario_cell}="Todos",'
            f'RSQ(FILTER({enpi_rng},{valid_all}),FILTER({ref_rng},{valid_all})),'
            f'RSQ(FILTER({enpi_rng},{valid_scen}),FILTER({ref_rng},{valid_scen}))),"")'
        )
    raise ValueError(metric)


def main() -> None:
    wb = load_workbook(INPUT)
    ws = wb["Todos"]
    first_row = 2
    last_row = ws.max_row

    if "Indicadores_Formulas" in wb.sheetnames:
        del wb["Indicadores_Formulas"]
    if "Notas_Calculo" in wb.sheetnames:
        del wb["Notas_Calculo"]

    helper_headers = {
        "AC": "Corriente Metrel A",
        "AD": "Potencia Metrel kW derivada",
        "AE": "Error Potencia kW ENPI-Metrel",
        "AF": "Reserva",
        "AG": "Reserva",
        "AH": "Reserva",
        "AI": "Reserva",
        "AJ": "Reserva",
        "AK": "Reserva",
        "AL": "Reserva",
        "AM": "Reserva",
        "AN": "Reserva",
        "AO": "Reserva",
        "AP": "Reserva",
        "AQ": "Reserva",
        "AR": "Reserva",
        "AS": "Reserva",
        "AT": "Reserva",
        "AU": "Error Voltaje ENPI-Metrel",
        "AV": "Error Corriente ENPI-Metrel",
        "AW": "Error Frecuencia ENPI-Metrel",
        "AX": "Error FP ENPI-Metrel",
        "AY": "Error Potencia ENPI-Metrel",
        "AZ": "Abs Error Voltaje",
        "BA": "Abs Error Corriente",
        "BB": "Abs Error Frecuencia",
        "BC": "Abs Error FP",
        "BD": "Abs Error Potencia",
        "BE": "APE Voltaje %",
        "BF": "APE Corriente %",
        "BG": "APE Frecuencia %",
        "BH": "APE FP %",
        "BI": "APE Potencia %",
        "BJ": "Sq Error Voltaje",
        "BK": "Sq Error Corriente",
        "BL": "Sq Error Frecuencia",
        "BM": "Sq Error FP",
        "BN": "Sq Error Potencia",
    }

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for col, header in helper_headers.items():
        cell = ws[f"{col}1"]
        cell.value = header
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[col].width = 16

    for row in range(first_row, last_row + 1):
        ws[f"AC{row}"] = f'=IF(OR(D{row}="003",D{row}="007"),K{row},K{row}/1000)'
        ws[f"AD{row}"] = f'=IFERROR(G{row}*AC{row}*Q{row}/1000,"")'
        ws[f"AE{row}"] = f'=IFERROR(T{row}-AD{row},"")'
        ws[f"AU{row}"] = f'=IFERROR(R{row}-G{row},"")'
        ws[f"AV{row}"] = f'=IFERROR(S{row}-AC{row},"")'
        ws[f"AW{row}"] = f'=IFERROR(V{row}-O{row},"")'
        ws[f"AX{row}"] = f'=IFERROR(W{row}-Q{row},"")'
        ws[f"AY{row}"] = f'=IFERROR(T{row}-AD{row},"")'
        ws[f"AZ{row}"] = f'=IFERROR(ABS(AU{row}),"")'
        ws[f"BA{row}"] = f'=IFERROR(ABS(AV{row}),"")'
        ws[f"BB{row}"] = f'=IFERROR(ABS(AW{row}),"")'
        ws[f"BC{row}"] = f'=IFERROR(ABS(AX{row}),"")'
        ws[f"BD{row}"] = f'=IFERROR(ABS(AY{row}),"")'
        ws[f"BE{row}"] = f'=IFERROR(ABS(AU{row}/G{row})*100,"")'
        ws[f"BF{row}"] = f'=IFERROR(ABS(AV{row}/AC{row})*100,"")'
        ws[f"BG{row}"] = f'=IFERROR(ABS(AW{row}/O{row})*100,"")'
        ws[f"BH{row}"] = f'=IFERROR(ABS(AX{row}/Q{row})*100,"")'
        ws[f"BI{row}"] = f'=IFERROR(ABS(AY{row}/AD{row})*100,"")'
        ws[f"BJ{row}"] = f'=IFERROR(AU{row}^2,"")'
        ws[f"BK{row}"] = f'=IFERROR(AV{row}^2,"")'
        ws[f"BL{row}"] = f'=IFERROR(AW{row}^2,"")'
        ws[f"BM{row}"] = f'=IFERROR(AX{row}^2,"")'
        ws[f"BN{row}"] = f'=IFERROR(AY{row}^2,"")'

    ind = wb.create_sheet("Indicadores_Formulas", 0)
    headers = [
        "Escenario",
        "Variable",
        "N",
        "MAE",
        "RMSE",
        "MAPE %",
        "Bias ENPI-Metrel",
        "Max Abs Error",
        "R2",
        "Unidad",
    ]
    ind.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ind[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    row_idx = 2
    for scenario in SCENARIOS:
        for var in VARIABLES:
            ind.cell(row_idx, 1).value = scenario
            ind.cell(row_idx, 2).value = var["name"]
            ind.cell(row_idx, 3).value = metric_formula("N", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 4).value = metric_formula("MAE", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 5).value = metric_formula("RMSE", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 6).value = metric_formula("MAPE", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 7).value = metric_formula("Bias", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 8).value = metric_formula("Max", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 9).value = metric_formula("R2", f"$A{row_idx}", var, first_row, last_row)
            ind.cell(row_idx, 10).value = var["unit"]
            row_idx += 1

    widths = [14, 24, 10, 14, 14, 14, 18, 16, 12, 10]
    for col, width in enumerate(widths, start=1):
        ind.column_dimensions[get_column_letter(col)].width = width
    for row in range(2, ind.max_row + 1):
        for col in range(4, 10):
            ind.cell(row, col).number_format = "0.000000"
    ind.freeze_panes = "A2"

    notes = wb.create_sheet("Notas_Calculo", 1)
    notes["A1"] = "Notas metodologicas"
    notes["A1"].font = Font(bold=True, size=14)
    notes["A3"] = "1. Todos los indicadores de la hoja Indicadores_Formulas son formulas de Excel, no valores pegados desde Python."
    notes["A4"] = "2. La corriente Metrel se normaliza a amperes en Todos!AC: escenarios 003 y 007 ya vienen en A; los demas vienen en mA."
    notes["A5"] = "3. La potencia Metrel es derivada: Voltaje Metrel * Corriente Metrel A * Factor de potencia Metrel / 1000."
    notes["A6"] = "4. Energia_kWh ENPI no tiene indicador de error porque el libro no incluye energia Metrel equivalente."
    notes["A7"] = "5. MAPE omite automaticamente filas sin referencia valida o con division invalida mediante IFERROR."
    notes["A8"] = "6. Los rangos usados son los de la hoja Todos corregida, despues de retirar perdidas de paquetes/tramos descartados."
    notes.column_dimensions["A"].width = 140

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
